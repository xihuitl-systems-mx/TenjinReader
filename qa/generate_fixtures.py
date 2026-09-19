from pathlib import Path

from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from reportlab.lib.colors import HexColor, white
from reportlab.pdfgen import canvas


OUTPUT_DIR = Path(__file__).resolve().parent
GREEN = RGBColor(0x23, 0x66, 0x54)
GOLD = RGBColor(0xEF, 0xB9, 0x33)


def make_pdf() -> None:
    path = OUTPUT_DIR / "sample.pdf"
    document = canvas.Canvas(str(path), pagesize=(612, 792))
    document.setTitle("Documento de prueba TenjinReader")
    document.setAuthor("TenjinReader QA")

    document.setFillColor(HexColor("#236654"))
    document.rect(0, 700, 612, 92, stroke=0, fill=1)
    document.setFillColor(white)
    document.setFont("Helvetica-Bold", 28)
    document.drawString(52, 742, "TenjinReader")
    document.setFont("Helvetica", 12)
    document.drawString(52, 718, "Documento PDF de prueba")
    document.setFillColor(HexColor("#171a18"))
    document.setFont("Helvetica-Bold", 20)
    document.drawString(52, 650, "Lectura y edición local")
    document.setFont("Helvetica", 13)
    document.drawString(52, 616, "Busca la palabra ligereza y añade una anotación.")
    document.drawString(52, 592, "Este archivo tiene dos páginas y un enlace seguro.")
    document.setFillColor(HexColor("#efb933"))
    document.roundRect(52, 500, 240, 56, 8, stroke=0, fill=1)
    document.setFillColor(HexColor("#174c3e"))
    document.setFont("Helvetica-Bold", 14)
    document.drawCentredString(172, 521, "Ligereza sin nube")
    document.setFillColor(HexColor("#236654"))
    document.setFont("Helvetica", 12)
    document.drawString(52, 452, "Documentación de PDF.js")
    document.linkURL(
        "https://mozilla.github.io/pdf.js/",
        (50, 444, 220, 468),
        relative=0,
        thickness=0,
    )
    document.showPage()

    document.setFillColor(HexColor("#f3f5f1"))
    document.rect(0, 0, 612, 792, stroke=0, fill=1)
    document.setFillColor(HexColor("#236654"))
    document.setFont("Helvetica-Bold", 26)
    document.drawString(52, 700, "Página 2")
    document.setFillColor(HexColor("#171a18"))
    document.setFont("Helvetica", 14)
    document.drawString(52, 660, "Usa las flechas para volver y prueba el zoom.")
    document.setFillColor(HexColor("#efb933"))
    document.circle(120, 500, 58, stroke=0, fill=1)
    document.setFillColor(HexColor("#236654"))
    document.rect(210, 450, 250, 100, stroke=0, fill=1)
    document.setFillColor(white)
    document.setFont("Helvetica-Bold", 18)
    document.drawCentredString(335, 492, "Todo queda en tu equipo")
    document.save()


def add_textbox(slide, text, left, top, width, height, size, color, bold=False):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.clear()
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.LEFT
    run = paragraph.add_run()
    run.text = text
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def make_pptx() -> None:
    path = OUTPUT_DIR / "sample.pptx"
    deck = Presentation()
    deck.slide_width = Inches(13.333333)
    deck.slide_height = Inches(7.5)

    slide = deck.slides.add_slide(deck.slide_layouts[6])
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = RGBColor(0xF3, 0xF5, 0xF1)
    accent = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.34), Inches(7.5)
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = GREEN
    accent.line.fill.background()
    add_textbox(
        slide,
        "PLUMA READER",
        Inches(0.9),
        Inches(0.8),
        Inches(4),
        Inches(0.5),
        14,
        GREEN,
        True,
    )
    add_textbox(
        slide,
        "Presentaciones,\nsin peso extra.",
        Inches(0.9),
        Inches(1.55),
        Inches(6.3),
        Inches(2.1),
        38,
        RGBColor(0x17, 0x1A, 0x18),
        True,
    )
    add_textbox(
        slide,
        "PPTX se abre en modo de solo lectura y se procesa localmente.",
        Inches(0.9),
        Inches(4.0),
        Inches(5.5),
        Inches(0.8),
        18,
        RGBColor(0x52, 0x59, 0x55),
    )
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(8.0),
        Inches(1.25),
        Inches(4.1),
        Inches(4.7),
    )
    card.fill.solid()
    card.fill.fore_color.rgb = GREEN
    card.line.fill.background()
    dot = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(9.25), Inches(2.05), Inches(1.6), Inches(1.6)
    )
    dot.fill.solid()
    dot.fill.fore_color.rgb = GOLD
    dot.line.fill.background()
    add_textbox(
        slide,
        "PDF + PPTX",
        Inches(8.9),
        Inches(4.15),
        Inches(2.3),
        Inches(0.5),
        21,
        RGBColor(0xFF, 0xFF, 0xFF),
        True,
    )

    slide = deck.slides.add_slide(deck.slide_layouts[6])
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    add_textbox(
        slide,
        "Rendimiento local",
        Inches(0.8),
        Inches(0.55),
        Inches(6),
        Inches(0.7),
        30,
        GREEN,
        True,
    )
    add_textbox(
        slide,
        "Una gráfica sencilla para comprobar el renderizado del visor.",
        Inches(0.82),
        Inches(1.25),
        Inches(7),
        Inches(0.5),
        15,
        RGBColor(0x52, 0x59, 0x55),
    )
    data = ChartData()
    data.categories = ["PDF", "PPTX", "Nube"]
    data.add_series("Prioridad", (10, 8, 0))
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(1.0),
        Inches(2.0),
        Inches(7.2),
        Inches(4.3),
        data,
    ).chart
    chart.has_legend = False
    chart.value_axis.maximum_scale = 10
    chart.value_axis.minimum_scale = 0
    chart.series[0].format.fill.solid()
    chart.series[0].format.fill.fore_color.rgb = GREEN
    callout = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(9.0),
        Inches(2.25),
        Inches(3.2),
        Inches(2.9),
    )
    callout.fill.solid()
    callout.fill.fore_color.rgb = GOLD
    callout.line.fill.background()
    add_textbox(
        slide,
        "Ligero\nRápido\nPrivado",
        Inches(9.45),
        Inches(2.8),
        Inches(2.3),
        Inches(1.8),
        25,
        GREEN,
        True,
    )

    # Remove the empty default slide if a template ever starts with one.
    while len(deck.slides) > 2:
        slide_id = deck.slides._sldIdLst[0]
        relationship_id = slide_id.rId
        deck.part.drop_rel(relationship_id)
        del deck.slides._sldIdLst[0]
    deck.save(path)


if __name__ == "__main__":
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    make_pdf()
    make_pptx()
